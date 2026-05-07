import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_excel_bytes(data: list[dict], columns: list[str]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scraped Data"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    alt_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    for row_idx, row in enumerate(data, start=2):
        for col_idx, col_name in enumerate(columns, start=1):
            value = row.get(col_name) if isinstance(row, dict) else None
            if value is None:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if not isinstance(value, (int, float)) else value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    ws.row_dimensions[1].height = 30
    for col_idx, col_name in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row in data:
            val = str(row.get(col_name, "") if isinstance(row, dict) else "")
            max_len = max(max_len, min(len(val), 60))
        ws.column_dimensions[col_letter].width = max(12, max_len + 4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
